import io

import docx
import pymupdf


def extract_pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc).strip()


def extract_docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)